#!/usr/bin/env python3
"""Génère une liste de lots simple (un palier de prix) à partir de fichiers de recherche JSON.

Usage :
    python scripts/build_liste.py --research <dossier_json> --out <fichier.xlsx> --total 100 \
        --title "100 nouveaux vinyles & CD" --price-min 45 --price-max 60 \
        [--quota Vinyle=50 --quota CD=50] [--exclude data/lots_musique.json] [--data-out data/x.json]

L'ordre des fichiers de recherche est conservé (les recherches classent le plus connu en premier).
Les quotas sont appliqués catégorie par catégorie, puis la liste est complétée jusqu'à --total.
Le surplus va dans l'onglet « Réserve ».
"""
import argparse
import glob
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Border, Font

from build_lots_musique import FIELDS, FONT, THIN, norm_url, style_header, write_items


def load(research_dir, pmin, pmax, excluded):
    rows, seen = [], set(excluded)
    for path in sorted(glob.glob(os.path.join(research_dir, "*.json"))):
        with open(path, encoding="utf-8") as f:
            for it in json.load(f):
                row = {k: it.get(k, "") for k in FIELDS}
                try:
                    row["price_myr"] = round(float(row["price_myr"]), 2)
                except (TypeError, ValueError):
                    continue
                key = norm_url(row["url"]) if row["url"] else None
                if not key or key in seen or not pmin <= row["price_myr"] <= pmax:
                    continue
                seen.add(key)
                rows.append(row)
    return rows


def select(rows, total, quotas):
    chosen = []
    for cat, n in quotas.items():
        chosen += [r for r in rows if r["category"] == cat][:n]
    picked = {id(r) for r in chosen}
    # Complète en alternant les catégories pour garder le mélange
    rest = [r for r in rows if id(r) not in picked]
    cats = list(dict.fromkeys(r["category"] for r in rest))
    queues = {c: [r for r in rest if r["category"] == c] for c in cats}
    while len(chosen) < total and any(queues.values()):
        for c in cats:
            if queues[c] and len(chosen) < total:
                chosen.append(queues[c].pop(0))
    picked = {id(r) for r in chosen}
    order = list(dict.fromkeys(r["category"] for r in rows))
    chosen.sort(key=lambda r: order.index(r["category"]))
    return chosen, [r for r in rows if id(r) not in picked]


def write_summary(ws, title, subtitle, cats, n_lots, n_reserve):
    ws["A1"] = title
    ws["A1"].font = Font(name=FONT, size=14, bold=True)
    ws["A2"] = subtitle
    ws["A2"].font = Font(name=FONT, italic=True, size=10)
    head = ["Catégorie", "Nb lots", "Prix moyen (RM)", "Budget (RM)", "En réserve"]
    for c, h in enumerate(head, start=1):
        ws.cell(row=4, column=c, value=h)
    style_header(ws, 4, len(head))
    last, rlast = n_lots + 1, max(n_reserve + 1, 2)
    first = r = 5
    for cat in cats:
        ws.cell(row=r, column=1, value=cat)
        ws.cell(row=r, column=2, value=f"=COUNTIF(Lots!$B$2:$B${last},A{r})")
        ws.cell(row=r, column=3, value=f"=IFERROR(AVERAGEIF(Lots!$B$2:$B${last},A{r},Lots!$D$2:$D${last}),0)")
        ws.cell(row=r, column=4, value=f"=SUMIF(Lots!$B$2:$B${last},A{r},Lots!$D$2:$D${last})")
        ws.cell(row=r, column=5, value=f"=COUNTIF(Réserve!$B$2:$B${rlast},A{r})")
        r += 1
    ws.cell(row=r, column=1, value="Total")
    ws.cell(row=r, column=2, value=f"=SUM(B{first}:B{r - 1})")
    ws.cell(row=r, column=3, value=f"=IFERROR(AVERAGE(Lots!$D$2:$D${last}),0)")
    ws.cell(row=r, column=4, value=f"=SUM(D{first}:D{r - 1})")
    ws.cell(row=r, column=5, value=f"=SUM(E{first}:E{r - 1})")
    for row in ws.iter_rows(min_row=first, max_row=r, max_col=len(head)):
        for cell in row:
            cell.font = Font(name=FONT, bold=cell.row == r)
            cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
        row[2].number_format = row[3].number_format = '"RM "#,##0.00'
    notes = [
        "Méthode",
        "• Prix = prix article affiché sur la boutique, hors frais de port.",
        "• Stock : disponibilité lue sur la boutique (données produit ou fiche), pas de précommande.",
        "• Onglet « Réserve » : produits valides en plus, pour remplacer un article épuisé.",
        "• Les stocks bougent vite (souvent 1–3 exemplaires) : revérifier au moment de commander.",
    ]
    for i, t in enumerate(notes):
        ws.cell(row=r + 2 + i, column=1, value=t).font = Font(name=FONT, bold=i == 0, size=10)
    for col, w in zip("ABCDE", [16, 10, 16, 14, 12]):
        ws.column_dimensions[col].width = w


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--research", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--total", type=int, required=True)
    ap.add_argument("--title", required=True)
    ap.add_argument("--price-min", type=float, required=True)
    ap.add_argument("--price-max", type=float, required=True)
    ap.add_argument("--quota", action="append", default=[], help="Catégorie=N")
    ap.add_argument("--exclude", action="append", default=[], help="JSON {lots, reserve} déjà utilisés")
    ap.add_argument("--data-out")
    ap.add_argument("--date", default="20/10/2026")
    args = ap.parse_args()

    excluded = set()
    for path in args.exclude:
        with open(path, encoding="utf-8") as f:
            d = json.load(f)
        excluded |= {norm_url(r["url"]) for r in d.get("lots", []) + d.get("reserve", [])}
    quotas = {k: int(v) for k, v in (q.split("=", 1) for q in args.quota)}

    rows = load(args.research, args.price_min, args.price_max, excluded)
    selected, reserve = select(rows, args.total, quotas)
    cats = list(dict.fromkeys(r["category"] for r in rows))

    wb = Workbook()
    ws = wb.active
    ws.title = "Résumé"
    write_items(wb.create_sheet("Lots"), selected, "Lots")
    write_items(wb.create_sheet("Réserve"), reserve, "Reserve")
    subtitle = (f"RM{args.price_min:g}–{args.price_max:g} · en stock · livrable en Malaisie · "
                f"relevé le {args.date}")
    write_summary(ws, args.title, subtitle, cats, len(selected), len(reserve))
    wb.save(args.out)
    if args.data_out:
        with open(args.data_out, "w", encoding="utf-8") as f:
            json.dump({"lots": selected, "reserve": reserve}, f, ensure_ascii=False, indent=1)

    counts = {c: sum(1 for r in selected if r["category"] == c) for c in cats}
    print(f"{len(rows)} produits valides -> {len(selected)} lots, {len(reserve)} en réserve : {counts}")


if __name__ == "__main__":
    main()
