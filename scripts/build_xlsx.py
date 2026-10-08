#!/usr/bin/env python3
"""Fusionne les fichiers de recherche JSON et génère l'Excel des leads
(magasins de vinyles + revendeurs de cartes Pokémon, Asie du Sud-Est).

Usage :
    python scripts/build_xlsx.py --research <dossier_json> --out <fichier.xlsx> [--data-dir data]

Chaque fichier JSON du dossier de recherche est un tableau d'objets avec les clés
de FIELDS ; les fichiers sont reconnus par leur préfixe (vinyl_*.json / pokemon_*.json).
"""
import argparse
import csv
import datetime
import glob
import json
import os
import re
import unicodedata
from urllib.parse import quote_plus, urlparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

FIELDS = [
    "name", "country", "city", "address", "type", "website", "instagram", "facebook",
    "tiktok", "phone_whatsapp", "email", "marketplaces", "google_maps", "specialty",
    "big_items", "size_tier", "live_selling", "status", "notes", "sources",
]

COUNTRY_FR = {
    "thailand": "Thaïlande", "thailande": "Thaïlande", "thaïlande": "Thaïlande",
    "vietnam": "Vietnam", "viet nam": "Vietnam",
    "philippines": "Philippines",
    "singapore": "Singapour", "singapour": "Singapour",
    "indonesia": "Indonésie", "indonésie": "Indonésie", "indonesie": "Indonésie",
}
COUNTRY_ORDER = ["Thaïlande", "Vietnam", "Philippines", "Singapour", "Indonésie"]

TYPE_FR = {
    "physical store": "Boutique physique",
    "physical + online": "Boutique + en ligne",
    "online only (website)": "En ligne (site web)",
    "marketplace seller": "Vendeur marketplace",
    "social media seller": "Vendeur réseaux sociaux",
    "market stall": "Stand / marché",
    "distributor / wholesaler": "Distributeur / grossiste",
    "café / bar with shop": "Café / bar avec boutique",
    "cafe / bar with shop": "Café / bar avec boutique",
    "chain / retailer": "Chaîne / grande enseigne",
}
PHYSICAL_TYPES = {
    "Boutique physique", "Boutique + en ligne", "Stand / marché",
    "Café / bar avec boutique", "Chaîne / grande enseigne",
}
STATUS_FR = {"active": "Actif", "unverified": "Non vérifié", "possibly closed": "Peut-être fermé"}
SIZE_FR = {"large": "Gros", "medium": "Moyen", "small": "Petit", "unknown": "Inconnu"}
SIZE_ORDER = {"Gros": 0, "Moyen": 1, "Petit": 2, "Inconnu": 3}

# Domaines partagés par plusieurs vendeurs : inutilisables pour dédoublonner.
SHARED_DOMAINS = (
    "facebook.com", "instagram.com", "linktr.ee", "shopee.", "lazada.", "tokopedia.com",
    "carousell.", "discogs.com", "tiktok.com", "google.", "goo.gl", "wa.me", "line.me",
    "bit.ly", "linkin.bio", "beacons.ai", "blibli.com", "bukalapak.com", "youtube.com",
    "lin.ee", "msha.ke", "taplink", "wanderlog.com", "tripadvisor.", "foursquare.com",
)

EMPTY_TOKENS = {"", "n/a", "na", "none", "null", "-", "unknown", "?", "nil", "tbd"}


def clean(value):
    if value is None:
        return ""
    if isinstance(value, list):
        return " | ".join(c for c in (clean(v) for v in value) if c)
    text = re.sub(r"\s+", " ", str(value)).strip()
    return "" if text.lower() in EMPTY_TOKENS else text


def norm_key(text):
    text = unicodedata.normalize("NFKD", text.lower())
    text = "".join(c for c in text if not unicodedata.combining(c))
    text = text.replace("đ", "d").replace("&", "and")
    return re.sub(r"[\W_]+", "", text)


def ig_handle(value):
    m = re.search(r"instagram\.com/([A-Za-z0-9_.]+)", value)
    if m:
        handle = m.group(1)
    else:
        m = re.match(r"^@?([A-Za-z0-9_.]{2,30})$", value)
        handle = m.group(1) if m else ""
    handle = handle.strip(".").lower()
    return "" if handle in {"p", "reel", "reels", "explore", "stories", "accounts"} else handle


def norm_instagram(value):
    value = clean(value)
    handle = ig_handle(value) if value else ""
    return f"https://www.instagram.com/{handle}" if handle else value


def norm_facebook(value):
    value = clean(value)
    if not value:
        return ""
    if re.match(r"^(www\.|m\.|web\.)?(facebook|fb)\.com/", value, re.I):
        value = "https://" + value
    if value.startswith("@"):
        return "https://www.facebook.com/" + value[1:]
    if value.lower().startswith("http"):
        value = re.sub(r"://(m\.|web\.)?(facebook|fb)\.com", "://www.facebook.com", value, flags=re.I)
        if "profile.php" not in value and "/groups/" not in value:
            value = value.split("?")[0]
        return value.rstrip("/")
    return value


def norm_tiktok(value):
    value = clean(value)
    if not value:
        return ""
    m = re.search(r"tiktok\.com/@([A-Za-z0-9_.]+)", value)
    handle = m.group(1) if m else (value[1:] if re.match(r"^@[A-Za-z0-9_.]+$", value) else "")
    return f"https://www.tiktok.com/@{handle}" if handle else value


def norm_url(value):
    value = clean(value)
    if value and not value.lower().startswith("http") and re.match(r"^[\w.-]+\.[a-z]{2,}(/.*)?$", value, re.I):
        value = "https://" + value
    return value


def web_domain(url):
    if not url.lower().startswith("http"):
        return ""
    netloc = urlparse(url).netloc.lower()
    netloc = netloc[4:] if netloc.startswith("www.") else netloc
    return "" if any(d in netloc for d in SHARED_DOMAINS) else netloc


def normalize(raw, category):
    rec = {k: clean(raw.get(k)) for k in FIELDS if k != "sources"}
    sources = raw.get("sources") or []
    if isinstance(sources, str):
        sources = re.split(r"\s*[|,;\n]\s*", sources)
    rec["sources"] = [s for s in (clean(x) for x in sources) if s]
    rec["country"] = COUNTRY_FR.get(rec["country"].lower(), rec["country"])
    rec["type"] = TYPE_FR.get(rec["type"].lower(), rec["type"])
    rec["status"] = STATUS_FR.get(rec["status"].lower(), rec["status"] or "Non vérifié")
    rec["website"] = norm_url(rec["website"])
    rec["instagram"] = norm_instagram(rec["instagram"])
    rec["facebook"] = norm_facebook(rec["facebook"])
    rec["tiktok"] = norm_tiktok(rec["tiktok"])
    rec["google_maps"] = norm_url(rec["google_maps"])
    live = rec["live_selling"].lower()
    rec["live_selling"] = "Oui" if live in {"yes", "oui", "y", "true"} else ""
    if category == "pokemon":
        rec["size_tier"] = SIZE_FR.get(rec["size_tier"].lower(), "Inconnu")
    else:
        rec["size_tier"] = rec["big_items"] = ""
    return rec


def filled(rec):
    return sum(1 for k in FIELDS if rec.get(k))


def merge_group(group):
    group = sorted(group, key=filled, reverse=True)
    out = dict(group[0])
    for rec in group[1:]:
        for k in FIELDS:
            if k == "sources":
                continue
            if not out.get(k) and rec.get(k):
                out[k] = rec[k]
            elif k in ("specialty", "big_items", "notes", "marketplaces") and rec.get(k) \
                    and rec[k].lower() not in out[k].lower():
                out[k] = f"{out[k]} | {rec[k]}"
    # Un signalement de fermeture l'emporte, puis une preuve d'activité récente.
    statuses = {rec["status"] for rec in group}
    for status in ("Peut-être fermé", "Actif"):
        if status in statuses:
            out["status"] = status
            break
    seen, sources = set(), []
    for rec in group:
        for s in rec["sources"]:
            if s not in seen:
                seen.add(s)
                sources.append(s)
    out["sources"] = sources
    return out


def dedupe(records):
    parent = list(range(len(records)))

    def find(i):
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    owner = {}
    for i, rec in enumerate(records):
        keys = [("name", rec["country"], norm_key(rec["name"]))]
        if "instagram.com/" in rec["instagram"]:
            keys.append(("ig", rec["country"], rec["instagram"].lower()))
        if rec["facebook"].lower().startswith("https://www.facebook.com/") and "profile.php" not in rec["facebook"]:
            keys.append(("fb", rec["country"], rec["facebook"].lower()))
        domain = web_domain(rec["website"])
        if domain:
            keys.append(("web", rec["country"], domain))
        for key in keys:
            if not key[2]:
                continue
            if key in owner:
                parent[find(i)] = find(owner[key])
            else:
                owner[key] = i
    groups = {}
    for i in range(len(records)):
        groups.setdefault(find(i), []).append(records[i])
    merged = []
    for group in groups.values():
        if len(group) > 1:
            print("  fusion:", " = ".join(sorted({g["name"] for g in group})))
        merged.append(merge_group(group))
    return merged


def load(research_dir, category):
    records = []
    for path in sorted(glob.glob(os.path.join(research_dir, f"{category}_*.json"))):
        with open(path, encoding="utf-8") as fh:
            data = json.load(fh)
        kept = [normalize(r, category) for r in data if clean(r.get("name"))]
        print(f"{os.path.basename(path)}: {len(kept)} entrées")
        records.extend(kept)
    records = [r for r in records if r["country"] in COUNTRY_ORDER]
    merged = dedupe(records)
    for rec in merged:
        if not rec["google_maps"] and rec["type"] in PHYSICAL_TYPES:
            query = " ".join(x for x in (rec["name"], rec["address"] or rec["city"], rec["country"]) if x)
            rec["google_maps"] = "https://www.google.com/maps/search/?api=1&query=" + quote_plus(query)
            rec["maps_is_search"] = True
    if category == "pokemon":
        merged.sort(key=lambda r: (COUNTRY_ORDER.index(r["country"]), SIZE_ORDER.get(r["size_tier"], 3),
                                   r["city"].lower(), r["name"].lower()))
    else:
        merged.sort(key=lambda r: (COUNTRY_ORDER.index(r["country"]), r["city"].lower(), r["name"].lower()))
    return merged


# ---------------------------------------------------------------- Excel ----

FONT = "Arial"
HEADER_FILL = PatternFill("solid", start_color="1F3864")
HEADER_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT, size=10)
LINK_FONT = Font(name=FONT, size=10, color="0563C1", underline="single")
TITLE_FONT = Font(name=FONT, size=14, bold=True, color="1F3864")
THIN = Side(style="thin", color="BFBFBF")

# (en-tête, clé, largeur, genre) ; genre : text | url | maps | sources | formula
VINYL_COLUMNS = [
    ("#", "_id", 6, "text"), ("Nom", "name", 32, "text"), ("Pays", "country", 12, "text"),
    ("Ville", "city", 16, "text"), ("Adresse / zone", "address", 38, "text"), ("Type", "type", 22, "text"),
    ("Spécialité", "specialty", 42, "text"), ("Site web", "website", 30, "url"),
    ("Instagram", "instagram", 30, "url"), ("Facebook", "facebook", 32, "url"), ("TikTok", "tiktok", 24, "url"),
    ("Téléphone / WhatsApp", "phone_whatsapp", 20, "text"), ("Email", "email", 26, "text"),
    ("Marketplaces (Shopee, Lazada, Discogs…)", "marketplaces", 32, "text"),
    ("Google Maps", "google_maps", 16, "maps"), ("Vente en live", "live_selling", 10, "text"),
    ("Statut", "status", 14, "text"), ("Nb canaux de contact", "_contacts", 11, "formula"),
    ("Notes", "notes", 45, "text"), ("Sources", "sources", 50, "sources"),
]
POKEMON_COLUMNS = [
    ("#", "_id", 6, "text"), ("Nom", "name", 32, "text"), ("Pays", "country", 12, "text"),
    ("Ville", "city", 16, "text"), ("Adresse / zone", "address", 38, "text"), ("Type", "type", 22, "text"),
    ("Taille reseller", "size_tier", 11, "text"), ("Gros items (preuves)", "big_items", 45, "text"),
    ("Priorité", "_priority", 9, "formula"), ("Spécialité (langues / produits)", "specialty", 40, "text"),
    ("Site web", "website", 30, "url"), ("Instagram", "instagram", 30, "url"), ("Facebook", "facebook", 32, "url"),
    ("TikTok", "tiktok", 24, "url"), ("Téléphone / WhatsApp", "phone_whatsapp", 20, "text"),
    ("Email", "email", 26, "text"), ("Marketplaces (Shopee, Lazada, Tokopedia…)", "marketplaces", 32, "text"),
    ("Google Maps", "google_maps", 16, "maps"), ("Vente en live", "live_selling", 10, "text"),
    ("Statut", "status", 14, "text"), ("Nb canaux de contact", "_contacts", 11, "formula"),
    ("Notes", "notes", 45, "text"), ("Sources", "sources", 50, "sources"),
]


def col_letter(columns, key):
    return get_column_letter([c[1] for c in columns].index(key) + 1)


def write_sheet(ws, columns, records, table_name):
    for idx, (header, _, width, _) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=idx, value=header)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[1].height = 30

    first_contact = col_letter(columns, "website")
    last_contact = col_letter(columns, "email")
    for r, rec in enumerate(records, start=2):
        for c, (_, key, _, kind) in enumerate(columns, start=1):
            cell = ws.cell(row=r, column=c)
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top")
            if key == "_id":
                cell.value = r - 1
            elif key == "_contacts":
                cell.value = f"=COUNTA({first_contact}{r}:{last_contact}{r})"
            elif key == "_priority":
                size, big, live = (col_letter(columns, k) for k in ("size_tier", "big_items", "live_selling"))
                cell.value = (f'=IF(OR({size}{r}="Gros",AND({big}{r}<>"",{live}{r}="Oui")),"A",'
                              f'IF(OR({size}{r}="Moyen",{big}{r}<>""),"B","C"))')
            elif kind == "sources":
                if rec["sources"]:
                    cell.value = " | ".join(rec["sources"])
                    if rec["sources"][0].startswith("http"):
                        cell.hyperlink = rec["sources"][0]
            else:
                value = rec.get(key) or None
                if kind == "maps" and value:
                    cell.value = "Rechercher sur Maps" if rec.get("maps_is_search") else "Voir sur Maps"
                    cell.hyperlink = value
                    cell.font = LINK_FONT
                    continue
                cell.value = value
                if kind == "url" and value and value.startswith("http"):
                    cell.hyperlink = value
                    cell.font = LINK_FONT

    last_row = max(len(records) + 1, 2)
    table = Table(displayName=table_name, ref=f"A1:{get_column_letter(len(columns))}{last_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "C2"


def write_summary(ws, vinyl_n, pokemon_n):
    ws.sheet_view.showGridLines = False
    ws["A1"] = "Leads Asie du Sud-Est — Magasins de vinyles & revendeurs de cartes Pokémon"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (f"Recherche web réalisée le {datetime.date(2026, 10, 8).strftime('%d/%m/%Y')} — "
                "Thaïlande, Vietnam, Philippines, Singapour, Indonésie")
    ws["A2"].font = Font(name=FONT, size=10, italic=True, color="595959")

    headers = ["Pays", "Vinyles", "Vinyles statut « Actif »", "Vinyles avec ≥1 contact",
               "Pokémon", "Pokémon priorité A", "Pokémon priorité B", "Pokémon avec ≥1 contact"]
    top = 4
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=top, column=c, value=h)
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
    ws.row_dimensions[top].height = 32
    v, p = "Vinyles", "'Pokémon'"
    vc, vs, vk = (col_letter(VINYL_COLUMNS, k) for k in ("country", "status", "_contacts"))
    pc, pp, pk = (col_letter(POKEMON_COLUMNS, k) for k in ("country", "_priority", "_contacts"))
    for i, country in enumerate(COUNTRY_ORDER):
        r = top + 1 + i
        a = f"$A{r}"
        row = [
            country,
            f"=COUNTIF({v}!${vc}:${vc},{a})",
            f'=COUNTIFS({v}!${vc}:${vc},{a},{v}!${vs}:${vs},"Actif")',
            f'=COUNTIFS({v}!${vc}:${vc},{a},{v}!${vk}:${vk},">0")',
            f"=COUNTIF({p}!${pc}:${pc},{a})",
            f'=COUNTIFS({p}!${pc}:${pc},{a},{p}!${pp}:${pp},"A")',
            f'=COUNTIFS({p}!${pc}:${pc},{a},{p}!${pp}:${pp},"B")',
            f'=COUNTIFS({p}!${pc}:${pc},{a},{p}!${pk}:${pk},">0")',
        ]
        for c, val in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=val)
            cell.font = BODY_FONT
            cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
            cell.alignment = Alignment(horizontal="left" if c == 1 else "center")
    total_row = top + 1 + len(COUNTRY_ORDER)
    ws.cell(row=total_row, column=1, value="Total")
    for c in range(2, len(headers) + 1):
        letter = get_column_letter(c)
        ws.cell(row=total_row, column=c, value=f"=SUM({letter}{top + 1}:{letter}{total_row - 1})")
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=total_row, column=c)
        cell.font = Font(name=FONT, size=10, bold=True)
        cell.fill = PatternFill("solid", start_color="D9E1F2")
        cell.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
        cell.alignment = Alignment(horizontal="left" if c == 1 else "center")

    notes_top = total_row + 2
    notes = [
        ("Comment lire le fichier", True),
        ("• Onglet « Vinyles » : disquaires physiques, stands de marché, cafés/bars et boutiques hi-fi qui vendent "
         "des vinyles, vendeurs en ligne (site, Shopee/Lazada/Tokopedia, Instagram/Facebook, Discogs), labels/distributeurs.", False),
        ("• Onglet « Pokémon » : boutiques TCG et revendeurs de cartes Pokémon, triés par pays puis par taille. "
         "« Gros items (preuves) » = ce qui a été vu dans les sources (booster boxes, cases, ETB, slabs PSA/BGS/CGC, vintage…).", False),
        ("• Taille reseller : Gros = gros revendeur / distributeur / chaîne / grosse boutique en ligne ; Moyen = boutique "
         "dédiée vendant des boîtes scellées ; Petit = petit vendeur ; Inconnu = pas assez d'info.", False),
        ("• Priorité (formule) : A = taille « Gros », ou gros items prouvés + vente en live ; B = taille « Moyen » ou "
         "gros items prouvés ; C = le reste.", False),
        ("• Statut : Actif = activité vue en 2024-2026 ; Non vérifié = pas de preuve récente ; Peut-être fermé = une "
         "source évoque une fermeture (à vérifier avant contact).", False),
        ("• Nb canaux de contact (formule) = nombre de champs remplis parmi Site web, Instagram, Facebook, TikTok, "
         "Téléphone, Email.", False),
        ("• Google Maps : « Voir sur Maps » = lien trouvé dans les sources ; « Rechercher sur Maps » = lien de recherche "
         "généré à partir du nom et de l'adresse (boutiques physiques uniquement).", False),
        ("Méthode et limites", True),
        ("• Données collectées par recherche web (guides, annuaires, articles, forums, profils Instagram/Facebook/TikTok, "
         "boutiques Shopee/Lazada/Tokopedia/Carousell/Discogs), en anglais et en langues locales.", False),
        ("• Chaque ligne cite au moins une source (colonne « Sources », la première est cliquable). Aucun contact n'a été "
         "inventé : un champ vide signifie que l'information n'a pas été trouvée.", False),
        ("• Les petites boutiques changent vite (déménagements, fermetures, changement de compte Instagram) : vérifier "
         "le compte avant toute prise de contact.", False),
        ("• Doublons fusionnés automatiquement (même nom, même Instagram, même page Facebook ou même site web dans un pays).", False),
    ]
    for i, (text, bold) in enumerate(notes):
        cell = ws.cell(row=notes_top + i, column=1, value=text)
        cell.font = Font(name=FONT, size=10, bold=bold, color="1F3864" if bold else "000000")
    ws.column_dimensions["A"].width = 16
    for c in range(2, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 15


def export_data(data_dir, name, records, columns):
    os.makedirs(data_dir, exist_ok=True)
    clean_records = [{k: v for k, v in r.items() if k != "maps_is_search"} for r in records]
    with open(os.path.join(data_dir, f"{name}.json"), "w", encoding="utf-8") as fh:
        json.dump(clean_records, fh, ensure_ascii=False, indent=1)
    keys = [(h, k) for h, k, _, kind in columns if kind != "formula" and k != "_id"]
    with open(os.path.join(data_dir, f"{name}.csv"), "w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow([h for h, _ in keys])
        for r in clean_records:
            writer.writerow([" | ".join(r[k]) if k == "sources" else r.get(k, "") for _, k in keys])


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--research", required=True)
    parser.add_argument("--out", required=True)
    parser.add_argument("--data-dir", default="data")
    args = parser.parse_args()

    vinyl = load(args.research, "vinyl")
    pokemon = load(args.research, "pokemon")
    print(f"Après fusion : {len(vinyl)} vinyles, {len(pokemon)} Pokémon")
    for label, recs in (("vinyles", vinyl), ("pokemon", pokemon)):
        by_country = {c: sum(1 for r in recs if r["country"] == c) for c in COUNTRY_ORDER}
        print(f"  {label}: {by_country}")

    wb = Workbook()
    summary = wb.active
    summary.title = "Résumé"
    write_sheet(wb.create_sheet("Vinyles"), VINYL_COLUMNS, vinyl, "TblVinyles")
    write_sheet(wb.create_sheet("Pokémon"), POKEMON_COLUMNS, pokemon, "TblPokemon")
    write_summary(summary, len(vinyl), len(pokemon))
    wb.save(args.out)
    export_data(args.data_dir, "vinyles", vinyl, VINYL_COLUMNS)
    export_data(args.data_dir, "pokemon", pokemon, POKEMON_COLUMNS)
    print("Écrit :", args.out)


if __name__ == "__main__":
    main()
